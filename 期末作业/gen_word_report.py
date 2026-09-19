# -*- coding: utf-8 -*-
"""将 项目开发报告.md 转换为带截图的 Word 文档"""
import os
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
MD_PATH = os.path.join(BASE_DIR, "项目开发报告.md")
DOCX_PATH = os.path.join(BASE_DIR, "项目开发报告.docx")
SHOT_DIR = os.path.join(BASE_DIR, "screenshots")


def set_cell_bg(cell, color_hex):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    tc_pr.append(shd)


def add_code_block(doc, code_text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(code_text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x37, 0x47, 0x4F)
    # 灰色底纹
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "F5F5F5")
    pPr.append(shd)


def parse_table(lines):
    """解析 markdown 表格，返回 (header, rows)"""
    rows = []
    for line in lines:
        line = line.strip()
        if not line.startswith("|"):
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        rows.append(cells)
    if len(rows) < 2:
        return None, []
    header = rows[0]
    # rows[1] 是分隔线 |---|---|
    data_rows = rows[2:] if len(rows) > 2 else []
    return header, data_rows


def add_table(doc, header, data_rows):
    table = doc.add_table(rows=1 + len(data_rows), cols=len(header))
    table.style = "Table Grid"
    # 表头
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(header):
        hdr_cells[i].text = h
        for p in hdr_cells[i].paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(10)
        set_cell_bg(hdr_cells[i], "E65100")
        for p in hdr_cells[i].paragraphs:
            for r in p.runs:
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    # 数据行
    for ri, row in enumerate(data_rows):
        cells = table.rows[ri + 1].cells
        for ci, val in enumerate(row):
            if ci < len(cells):
                cells[ci].text = val
                for p in cells[ci].paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(9)


def find_image(img_ref):
    """根据图片引用找到实际文件路径"""
    name = os.path.basename(img_ref)
    candidates = [
        os.path.join(SHOT_DIR, name),
        os.path.join(BASE_DIR, img_ref),
    ]
    for c in candidates:
        if os.path.exists(c):
            return c
    if os.path.isdir(SHOT_DIR):
        for f in os.listdir(SHOT_DIR):
            if f.lower().endswith((".png", ".jpg", ".jpeg")):
                base = os.path.splitext(f)[0].lower()
                name_base = os.path.splitext(name)[0].lower()
                if base == name_base or base.endswith(name_base) or name_base.endswith(base):
                    return os.path.join(SHOT_DIR, f)
    return None


def extract_screenshot_paths(text):
    """从文本中提取所有 screenshots/*.png 或 screenshots/*.jpg 路径"""
    return re.findall(r"screenshots/[\w\-]+\.(?:png|jpg|jpeg)", text, re.IGNORECASE)


def embed_images_from_text(doc, text):
    """从文本中提取截图路径并嵌入图片，返回去掉截图路径后的文本"""
    paths = extract_screenshot_paths(text)
    if not paths:
        return text
    for p in paths:
        img_path = find_image(p)
        if img_path:
            try:
                doc.add_picture(img_path, width=Inches(5.5))
                last_par = doc.paragraphs[-1]
                last_par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            except Exception:
                pass
    # 从文本中移除截图路径
    clean = re.sub(r"[（(]?\s*截图[：:]?\s*screenshots/[\w\-]+\.(?:png|jpg|jpeg)\s*[)）]?", "", text)
    clean = re.sub(r"screenshots/[\w\-]+\.(?:png|jpg|jpeg)", "", clean)
    return clean.strip()


def main():
    with open(MD_PATH, "r", encoding="utf-8") as f:
        content = f.read()

    lines = content.split("\n")
    doc = Document()

    # 设置默认字体
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

    i = 0
    in_code = False
    code_buf = []
    table_buf = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # 代码块
        if stripped.startswith("```"):
            if in_code:
                add_code_block(doc, "\n".join(code_buf))
                code_buf = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue

        # 表格收集
        if stripped.startswith("|") and "|" in stripped[1:]:
            table_buf.append(line)
            i += 1
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_buf.append(lines[i])
                i += 1
            header, data_rows = parse_table(table_buf)
            if header:
                add_table(doc, header, data_rows)
                # 提取表格中所有截图路径并嵌入
                all_cell_text = " ".join(header + [c for r in data_rows for c in r])
                shot_paths = extract_screenshot_paths(all_cell_text)
                for sp in shot_paths:
                    img_path = find_image(sp)
                    if img_path:
                        try:
                            doc.add_picture(img_path, width=Inches(5.5))
                            last_par = doc.paragraphs[-1]
                            last_par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        except Exception:
                            pass
            table_buf = []
            doc.add_paragraph()
            continue

        # 标题
        if stripped.startswith("# "):
            h = doc.add_heading(stripped[2:].strip(), level=0)
            for r in h.runs:
                r.font.color.rgb = RGBColor(0xBF, 0x36, 0x0C)
        elif stripped.startswith("## "):
            h = doc.add_heading(stripped[3:].strip(), level=1)
            for r in h.runs:
                r.font.color.rgb = RGBColor(0xE6, 0x51, 0x00)
        elif stripped.startswith("### "):
            h = doc.add_heading(stripped[4:].strip(), level=2)
            for r in h.runs:
                r.font.color.rgb = RGBColor(0x5D, 0x40, 0x37)
        elif stripped.startswith("#### "):
            doc.add_heading(stripped[5:].strip(), level=3)
        elif stripped.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            run = p.add_run(stripped[2:].strip())
            run.font.color.rgb = RGBColor(0x8D, 0x6E, 0x63)
            run.italic = True
        elif stripped.startswith("- ") or stripped.startswith("* "):
            text = stripped[2:].strip()
            # 清理标记
            text = text.replace("`", "")
            text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
            # 嵌入截图
            text = embed_images_from_text(doc, text)
            if text:
                doc.add_paragraph(text, style="List Bullet")
        elif re.match(r"^\d+\.\s", stripped):
            text = re.sub(r"^\d+\.\s", "", stripped)
            doc.add_paragraph(text, style="List Number")
        elif stripped == "---":
            p = doc.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            pBdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:space"), "1")
            bottom.set(qn("w:color"), "E65100")
            pBdr.append(bottom)
            pPr.append(pBdr)
        elif stripped == "":
            pass  # 空行跳过
        else:
            # 普通段落，处理行内截图引用
            clean = stripped.replace("`", "")
            clean = re.sub(r"\*\*(.*?)\*\*", r"\1", clean)
            # 嵌入截图并清理路径文本
            clean = embed_images_from_text(doc, clean)
            if clean:
                p = doc.add_paragraph(clean)
                p.paragraph_format.first_line_indent = Cm(0.74)

        i += 1

    doc.save(DOCX_PATH)
    print(f"Word 文档已生成: {DOCX_PATH}")


if __name__ == "__main__":
    main()
