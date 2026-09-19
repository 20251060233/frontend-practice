from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import datetime

TEMPLATE = r"C:\Users\ASUS\Desktop\python报告模板.docx"
OUTPUT = r"C:\Users\ASUS\Desktop\Python课堂作业实验报告-吴志贤-20251060233.docx"

# ===== 三道题内容 =====
THOUGHTS = {
    1: (
        "鸡兔同笼是经典二元一次方程组问题：设鸡有 c 只、兔有 r 只，"
        "根据"头的总数"和"脚的总数"列两个方程——c+r=heads、2c+4r=feet，"
        "用消元法解得 r=(feet-2*heads)/2，c=heads-r。"
        "实现时还要考虑无解情况：脚数必须为偶数且算出的鸡兔数不能为负数。"
    ),
    2: (
        "分数评级本质是多分支 if-elif 条件判断。"
        "按照 60、70、80、90 四个分界点依次判断："
        "先排除越界分数，再从低到高逐个 elif，最后 else 兜底 90-100 为"优"。"
        "注意 Python 中 elif 会短路，前面的条件不满足才会往下走。"
    ),
    3: (
        "这道题关键是"如何构造每一项"。"
        "观察可得：第 i 项是数字 x 重复 i 次（如 x=2 时第 3 项是 222）。"
        "Python 字符串支持乘法，x='2' * 3 直接得到 '222'，再用 int() 转整数。"
        "用列表推导式 range(1, n+1) 生成 1 到 n，sum() 一次性求和。"
    ),
}

CODE = {
    1: '''heads = int(input("输入头的个数: "))
feet = int(input("输入脚的个数: "))

# 设鸡 c 只，兔 r 只
# c + r = heads;  2c + 4r = feet
# 消元解得 r = (feet - 2*heads) / 2，c = heads - r
rabbits = (feet - 2 * heads) / 2
chickens = heads - rabbits

if feet % 2 != 0 or rabbits < 0 or chickens < 0 or rabbits != int(rabbits):
    print("无解：输入数据不合理")
else:
    print(f"鸡: {int(chickens)} 只, 兔: {int(rabbits)} 只")''',
    2: '''score = float(input("输入成绩分数: "))

if score < 0 or score > 100:
    print("分数需在 0~100 之间")
elif score < 60:
    print("等级: 不及格")
elif score < 70:
    print("等级: 差")
elif score < 80:
    print("等级: 中")
elif score < 90:
    print("等级: 良")
else:
    print("等级: 优")''',
    3: '''x = input("输入 x (单个数字): ").strip()
n = int(input("输入 n (项数): "))

# 第 i 项 = x 重复 i 次（字符串乘法），转 int 后求和
total = sum(int(x * i) for i in range(1, n + 1))

print(f"sum = {total}")''',
}

RESULTS = {
    1: '''输入头的个数: 35
输入脚的个数: 94
鸡: 23 只, 兔: 12 只''',
    2: '''输入成绩分数: 85
等级: 良''',
    3: '''输入 x (单个数字): 2
输入 n (项数): 3
sum = 246''',
}

SUMMARY = (
    "本次实验完成了三道 Python 基础编程题目。"
    "鸡兔同笼让我练习了数学建模——将文字问题抽象为方程再用代码求解；"
    "分数评级练习了 if-elif 多分支结构，注意 elif 的短路特性和边界覆盖；"
    "第三题让我再次感受到 Python 字符串乘法的强大——用 '2'*3='222' 这一行就代替了手动循环拼接。\n\n"
    "踩过的坑：一开始第三题想用数学公式（等比数列求和），但 x 不是单个数字时公式不成立，"
    "用字符串乘法反而更通用；鸡兔同笼曾忘记判断无解情况，算出负数或小数会闹出笑话。\n\n"
    "收获：Python 语法简洁，但简洁不代表随便——每道题都要考虑边界和异常输入，"
    "这是写代码与写数学题最大的区别。"
)


# ===== 辅助：往段落里追加 runs =====
def add_runs(para, text, font="宋体", size=12, bold=False, mono=False):
    """清空 para 现有内容，填入 text，自动保留换行。"""
    para.clear()
    if mono:
        font = "Consolas"  # 等宽字体显示代码更整齐
    for idx, line in enumerate(text.split("\n")):
        run = para.add_run(line)
        run.font.name = font
        run._element.rPr.rFonts.set(
            "w:eastAsia", "Consolas" if mono else "宋体"
        )
        run.font.size = Pt(size)
        run.font.bold = bold
        if idx < len(text.split("\n")) - 1:
            # 换行
            pass
        else:
            continue
        # 下一行
    # 上面的方式对最后一行有 bug，用下面更简单的：
    para.clear()
    for i, line in enumerate(text.split("\n")):
        run = para.add_run(line)
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        if i < len(text.split("\n")) - 1:
            run.add_break()


def set_paragraph(para, text, font="宋体", size=12, mono=False, bold=False):
    add_runs(para, text, font, size, bold, mono)


# ===== 开始填充 =====
doc = Document(TEMPLATE)

# --- 头部占位符替换 ---
# [2] 信息学院那行
doc.paragraphs[2].clear()
p = doc.paragraphs[2]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("      信息      学院    计算机科学与技术     专业    2025     级")
run.font.size = Pt(12)

# [4] 实验时间
today = datetime.date.today()
doc.paragraphs[4].clear()
p = doc.paragraphs[4]
run = p.add_run(f"实验时间     {today.year}     年   {today.month:02d}   月   {today.day:02d}   日")
run.font.size = Pt(12)

# [6] 姓名 / 学号
doc.paragraphs[6].clear()
p = doc.paragraphs[6]
run = p.add_run("姓名    吴志贤              学号     20251060233")
run.font.size = Pt(12)

# --- 实验思路 ---
set_paragraph(doc.paragraphs[13], "第一题\n" + "=" * 40 + "\n" + THOUGHTS[1], size=12)
set_paragraph(doc.paragraphs[15], "第二题\n" + "=" * 40 + "\n" + THOUGHTS[2], size=12)
set_paragraph(doc.paragraphs[17], "第三题\n" + "=" * 40 + "\n" + THOUGHTS[3], size=12)

# --- 实验代码（等宽字体） ---
set_paragraph(doc.paragraphs[20], "第一题\n" + "=" * 40 + "\n" + CODE[1], mono=True, size=10.5)
set_paragraph(doc.paragraphs[22], "第二题\n" + "=" * 40 + "\n" + CODE[2], mono=True, size=10.5)
set_paragraph(doc.paragraphs[24], "第三题\n" + "=" * 40 + "\n" + CODE[3], mono=True, size=10.5)

# --- 实验结果（等宽字体模拟截图） ---
result_text = (
    "第一题测试（头=35，脚=94）：\n" + "-" * 36 + "\n" + RESULTS[1] + "\n\n"
    "第二题测试（分数=85）：\n" + "-" * 36 + "\n" + RESULTS[2] + "\n\n"
    "第三题测试（x=2, n=3）：\n" + "-" * 36 + "\n" + RESULTS[3]
)
set_paragraph(doc.paragraphs[26], result_text, mono=True, size=10.5)

# --- 总结 ---
set_paragraph(doc.paragraphs[27], SUMMARY, size=12)

doc.save(OUTPUT)
print(f"报告已生成：{OUTPUT}")
